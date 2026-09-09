from __future__ import annotations

import io
import os
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("DATA_DIR", BASE_DIR / "data"))
DB_PATH = DATA_DIR / "consultas.db"
ONPE_URL = "https://consultaelectoral.onpe.gob.pe/"
MAX_RECORDS = 500

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024


def connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with connection() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS voters (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dni TEXT NOT NULL UNIQUE,
                status TEXT NOT NULL DEFAULT 'pendiente',
                region TEXT NOT NULL DEFAULT '',
                provincia TEXT NOT NULL DEFAULT '',
                distrito TEXT NOT NULL DEFAULT '',
                local_direccion TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL
            )
        """)


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_dni(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def valid_dni(value):
    return bool(re.fullmatch(r"\d{8}", value))


def row_dict(row):
    return dict(row)


@app.get("/")
def index():
    return render_template("index.html", onpe_url=ONPE_URL)


@app.get("/health")
def health():
    try:
        with connection() as conn:
            conn.execute("SELECT 1").fetchone()
        return jsonify(status="ok")
    except sqlite3.Error:
        return jsonify(status="error"), 503


@app.get("/api/records")
def records():
    with connection() as conn:
        rows = conn.execute("SELECT * FROM voters ORDER BY id DESC").fetchall()
    return jsonify(records=[row_dict(row) for row in rows])


@app.post("/api/records")
def create_record():
    payload = request.get_json(silent=True) or {}
    dni = clean_dni(payload.get("dni"))
    if not valid_dni(dni):
        return jsonify(error="El DNI debe contener exactamente 8 dígitos."), 400
    try:
        with connection() as conn:
            cursor = conn.execute(
                "INSERT INTO voters (dni, updated_at) VALUES (?, ?)", (dni, now_iso())
            )
            row = conn.execute("SELECT * FROM voters WHERE id = ?", (cursor.lastrowid,)).fetchone()
        return jsonify(record=row_dict(row)), 201
    except sqlite3.IntegrityError:
        return jsonify(error="Este DNI ya está en la lista."), 409


@app.post("/api/import")
def import_records():
    upload = request.files.get("file")
    if not upload or not upload.filename:
        return jsonify(error="Selecciona un archivo .xlsx."), 400
    if not upload.filename.lower().endswith(".xlsx"):
        return jsonify(error="Solo se admiten archivos Excel .xlsx."), 400
    try:
        workbook = load_workbook(upload.stream, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "").strip().upper() for cell in sheet[1]]
        if "DNI" not in headers:
            return jsonify(error="La primera fila debe incluir una columna llamada DNI."), 400
        dni_column = headers.index("DNI") + 1
        candidates = [clean_dni(sheet.cell(row=i, column=dni_column).value) for i in range(2, sheet.max_row + 1)]
    except Exception:
        return jsonify(error="No se pudo leer el Excel. Usa la plantilla disponible."), 400

    candidates = [dni for dni in candidates if dni]
    invalid = [dni for dni in candidates if not valid_dni(dni)]
    if invalid:
        preview = ", ".join(invalid[:3])
        return jsonify(error=f"Hay DNIs inválidos: {preview}. Cada DNI debe tener 8 dígitos."), 400
    unique = list(dict.fromkeys(candidates))
    if not unique:
        return jsonify(error="El archivo no contiene DNIs."), 400
    if len(unique) > MAX_RECORDS:
        return jsonify(error=f"El límite es de {MAX_RECORDS} DNIs por importación."), 400

    created = 0
    duplicates = 0
    with connection() as conn:
        for dni in unique:
            try:
                conn.execute("INSERT INTO voters (dni, updated_at) VALUES (?, ?)", (dni, now_iso()))
                created += 1
            except sqlite3.IntegrityError:
                duplicates += 1
    return jsonify(created=created, duplicates=duplicates)


@app.put("/api/records/<int:record_id>")
def update_record(record_id):
    payload = request.get_json(silent=True) or {}
    status = str(payload.get("status", "")).strip().lower()
    if status not in {"pendiente", "miembro", "no_miembro"}:
        return jsonify(error="Estado no válido."), 400
    fields = {
        "region": str(payload.get("region", "")).strip()[:80],
        "provincia": str(payload.get("provincia", "")).strip()[:80],
        "distrito": str(payload.get("distrito", "")).strip()[:80],
        "local_direccion": str(payload.get("local_direccion", "")).strip()[:240],
    }
    if status == "miembro" and not all(fields.values()):
        return jsonify(error="Completa región, provincia, distrito y dirección para un miembro de mesa."), 400
    with connection() as conn:
        existing = conn.execute("SELECT id FROM voters WHERE id = ?", (record_id,)).fetchone()
        if not existing:
            return jsonify(error="Registro no encontrado."), 404
        conn.execute("""
            UPDATE voters SET status=?, region=?, provincia=?, distrito=?, local_direccion=?, updated_at=?
            WHERE id=?
        """, (status, fields["region"], fields["provincia"], fields["distrito"], fields["local_direccion"], now_iso(), record_id))
        row = conn.execute("SELECT * FROM voters WHERE id = ?", (record_id,)).fetchone()
    return jsonify(record=row_dict(row))


@app.delete("/api/records/<int:record_id>")
def delete_record(record_id):
    with connection() as conn:
        cursor = conn.execute("DELETE FROM voters WHERE id = ?", (record_id,))
    if cursor.rowcount == 0:
        return jsonify(error="Registro no encontrado."), 404
    return "", 204


def styled_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Miembros de mesa"
    sheet.append(["DNI", "Región", "Provincia", "Distrito", "Dirección del local de votación"])
    for row in rows:
        sheet.append([row["dni"], row["region"], row["provincia"], row["distrito"], row["local_direccion"]])
    navy = "15253D"
    cream = "FFF9EF"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{max(sheet.max_row, 1)}"
    sheet.row_dimensions[1].height = 28
    widths = [14, 20, 22, 22, 48]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "@"
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=cream)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return workbook


@app.get("/api/template")
def template():
    return send_file(
        BASE_DIR / "plantilla_dni.xlsx",
        as_attachment=True,
        download_name="plantilla_dni.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/api/export")
def export():
    with connection() as conn:
        rows = conn.execute("SELECT * FROM voters WHERE status='miembro' ORDER BY dni").fetchall()
    workbook = styled_workbook(rows)
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name="miembros_de_mesa.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=False)

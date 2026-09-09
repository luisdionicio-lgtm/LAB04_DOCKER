import io

import pytest
from openpyxl import Workbook, load_workbook


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    import importlib
    import app as module
    importlib.reload(module)
    module.app.config.update(TESTING=True)
    return module.app.test_client()


def excel_bytes(values):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["DNI"])
    for value in values:
        sheet.append([value])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def test_health_and_page(client):
    assert client.get("/health").json == {"status": "ok"}
    assert b"Mesa Clara" in client.get("/").data


def test_create_validate_and_duplicate(client):
    assert client.post("/api/records", json={"dni": "123"}).status_code == 400
    assert client.post("/api/records", json={"dni": "12345678"}).status_code == 201
    assert client.post("/api/records", json={"dni": "12345678"}).status_code == 409


def test_import_update_and_export_only_members(client):
    response = client.post("/api/import", data={"file": (excel_bytes(["12345678", "87654321"]), "lista.xlsx")}, content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.json == {"created": 2, "duplicates": 0}
    records = client.get("/api/records").json["records"]
    member = next(item for item in records if item["dni"] == "12345678")
    response = client.put(f"/api/records/{member['id']}", json={"status": "miembro", "region": "Lima", "provincia": "Lima", "distrito": "Centro", "local_direccion": "Direcci\u00f3n de prueba"})
    assert response.status_code == 200
    export = client.get("/api/export")
    workbook = load_workbook(io.BytesIO(export.data), read_only=True)
    values = list(workbook.active.values)
    assert values[1][0] == "12345678"
    assert all("87654321" not in row for row in values)


def test_member_requires_location(client):
    record = client.post("/api/records", json={"dni": "12345678"}).json["record"]
    response = client.put(f"/api/records/{record['id']}", json={"status": "miembro"})
    assert response.status_code == 400

